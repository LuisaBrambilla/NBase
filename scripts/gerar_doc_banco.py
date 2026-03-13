"""
Gera um DOCX com a documentação do banco de dados e do processo de persistência das normas.
Saída: docs/Estrutura_Banco_e_Persistencia_Normas.docx
"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if bold:
        run.bold = True
    return p


def add_table_from_rows(doc, headers, rows, style="Table Grid"):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers), style=style)
    t.allow_autofit = False
    for j, h in enumerate(headers):
        t.rows[0].cells[j].text = h
        t.rows[0].cells[j].paragraphs[0].runs[0].bold = True
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            t.rows[i + 1].cells[j].text = str(cell)
    return t


def main():
    out_dir = ROOT / "docs"
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / "Estrutura_Banco_e_Persistencia_Normas.docx"

    doc = Document()
    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Documentação do Banco de Dados\nGen360 – Resoluções ANA / ARES-PCJ")
    r.bold = True
    r.font.size = Pt(16)
    doc.add_paragraph()
    doc.add_paragraph(
        "Este documento descreve a estrutura do banco de dados PostgreSQL utilizado no projeto, "
        "os campos das tabelas principais e o processo de desenvolvimento adotado para persistir "
        "os dados das normas (atos normativos) no banco."
    )
    doc.add_paragraph()

    # --- 1. Estrutura do banco ---
    add_heading(doc, "1. Estrutura do Banco de Dados", level=1)
    add_para(doc, "O banco utiliza PostgreSQL com as extensões uuid-ossp (geração de UUIDs) e pg_trgm (busca por similaridade de texto). As tabelas principais são norma (atos normativos) e admins (administradores do painel).", bold=False)
    doc.add_paragraph()

    add_heading(doc, "1.1 Tabela norma", level=2)
    add_para(doc, "Armazena os atos normativos (resoluções da ANA, ARES-PCJ, etc.). Chave natural: (orgao, tipo, ano, numero, origem).")
    doc.add_paragraph()

    headers_norma = ["Campo", "Tipo", "Restrições", "Descrição"]
    rows_norma = [
        ("id", "UUID", "PK, DEFAULT uuid_generate_v4()", "Identificador único do registro"),
        ("orgao", "TEXT", "NOT NULL, default 'ANA'", "Órgão emissor (ex.: ANA, ARES-PCJ)"),
        ("tipo", "TEXT", "NOT NULL, default 'Resolução'", "Tipo do ato (ex.: Resolução)"),
        ("numero", "INT", "NOT NULL", "Número do ato"),
        ("ano", "INT", "NOT NULL", "Ano do ato"),
        ("data_publicacao", "DATE", "", "Data de publicação"),
        ("titulo", "TEXT", "", "Título do ato"),
        ("ementa", "TEXT", "", "Ementa ou resumo"),
        ("assunto", "TEXT", "", "Assunto/classificação"),
        ("origem", "TEXT", "NOT NULL, CHECK", "Origem do dado: govbr, ocr, dou, dodf, manual, arespcj"),
        ("status", "TEXT", "NOT NULL, CHECK, default 'vigente'", "vigente | vigente_alterado | nao_vigente"),
        ("consolidado", "BOOLEAN", "NOT NULL, default FALSE", "Indica se é norma consolidada"),
        ("ignorar_estatisticas", "BOOLEAN", "NOT NULL, default FALSE", "Se TRUE, não entra nas estatísticas"),
        ("url_oficial", "TEXT", "", "URL do documento oficial na fonte"),
        ("caminho_html", "TEXT", "", "Caminho do HTML de conteúdo (ex.: data/html_conteudo/arespcj_2026_688.html)"),
        ("caminho_html_bruto", "TEXT", "", "Caminho do HTML completo (com cabeçalhos/rodapé)"),
        ("caminho_pdf", "TEXT", "", "Caminho do PDF, quando aplicável"),
        ("caminho_docx", "TEXT", "", "Nome do arquivo Word (.docx) gerado (ex.: arespcj_2026_688.docx)"),
        ("texto_busca", "TEXT", "", "Texto extraído para busca full-text (GIN com pg_trgm)"),
        ("criado_em", "TIMESTAMP", "DEFAULT now()", "Data/hora de criação do registro"),
        ("atualizado_em", "TIMESTAMP", "DEFAULT now()", "Atualizado pelo trigger antes de UPDATE"),
        ("verificacao_status", "TEXT", "", "Status da verificação por admin (preenchido no fluxo de verificação)"),
        ("verificacao_descricao", "TEXT", "", "Descrição/observação da verificação"),
        ("verificacao_em", "TIMESTAMP", "", "Data/hora em que o admin verificou"),
        ("verificacao_admin_id", "UUID", "", "ID do admin que realizou a verificação"),
    ]
    add_table_from_rows(doc, headers_norma, rows_norma)
    doc.add_paragraph()
    add_para(doc, "Constraint UNIQUE: (orgao, tipo, ano, numero, origem) — evita duplicatas por ato e origem.")
    add_para(doc, "Índices: ano+numero, data_publicacao, status, assunto; GIN em texto_busca (gin_trgm_ops); verificacao_status.")
    doc.add_paragraph()

    add_heading(doc, "1.2 Tabela admins", level=2)
    add_para(doc, "Administradores do painel. Cadastro apenas pelo backend (script criar_admin.py).")
    doc.add_paragraph()
    headers_admins = ["Campo", "Tipo", "Restrições", "Descrição"]
    rows_admins = [
        ("id", "UUID", "PK, DEFAULT uuid_generate_v4()", "Identificador único"),
        ("nome", "TEXT", "", "Nome do administrador"),
        ("email", "TEXT", "NOT NULL, UNIQUE", "Email de login"),
        ("senha_hash", "TEXT", "NOT NULL", "Hash da senha (PBKDF2)"),
        ("mfa_secret", "TEXT", "", "Segredo TOTP para MFA (2FA)"),
        ("mfa_enabled", "BOOLEAN", "NOT NULL, default TRUE", "Se MFA está ativo"),
        ("ativo", "BOOLEAN", "NOT NULL, default TRUE", "Se a conta está ativa"),
        ("criado_em", "TIMESTAMP", "DEFAULT now()", "Data de criação"),
        ("atualizado_em", "TIMESTAMP", "DEFAULT now()", "Atualizado por trigger"),
    ]
    add_table_from_rows(doc, headers_admins, rows_admins)
    doc.add_paragraph()

    add_heading(doc, "1.3 Triggers", level=2)
    add_para(doc, "A função atualizar_data() é disparada BEFORE UPDATE nas tabelas norma e admins, preenchendo automaticamente o campo atualizado_em com now().")
    doc.add_paragraph()

    # --- 2. Processo de persistência ---
    add_heading(doc, "2. Processo de Persistência dos Dados das Normas", level=1)
    add_para(doc, "No desenvolvimento, as normas são persistidas principalmente por scripts de coleta que obtêm dados de fontes externas, convertem documentos (PDF → HTML, PDF → Word) e gravam ou atualizam registros na tabela norma.")
    doc.add_paragraph()

    add_heading(doc, "2.1 Origens dos dados (origem)", level=2)
    add_para(doc, "O campo origem identifica a procedência do registro: govbr, ocr, dou, dodf, manual, arespcj. Para a ARES-PCJ utiliza-se origem = 'arespcj' e orgao = 'ARES-PCJ'. As migrações em db/migracoes/ estendem o CHECK de origem quando necessário (ex.: 005_origem_arespcj.sql, 006_caminho_docx.sql).")
    doc.add_paragraph()

    add_heading(doc, "2.2 Fluxo ARES-PCJ (coletor)", level=2)
    add_para(doc, "O script scripts/coletar_arespcj.py implementa o processo completo de persistência para as resoluções da ARES-PCJ:")
    doc.add_paragraph()
    steps = [
        "1. Acessar a página de listagem (URL_LISTAGEM) e extrair links dos PDFs e metadados (número, ano, data, título).",
        "2. Para cada resolução: baixar o PDF para data/pdf_arespcj/ (prefixo arespcj_<ano>_<numero>.pdf).",
        "3. Converter PDF → HTML com lib/pdf_para_html (texto, imagens, tabelas); salvar em data/html_conteudo/ e data/html_completo/.",
        "4. Converter PDF → Word com lib/pdf_para_docx (pdf2docx ou OCR para escaneados); salvar em data/docx_arespcj/ e guardar o nome do ficheiro em caminho_docx.",
        "5. Montar o dicionário de dados (orgao, tipo, numero, ano, data_publicacao, titulo, ementa, origem, status, url_oficial, caminho_html, caminho_html_bruto, caminho_docx, texto_busca) e chamar inserir_ou_atualizar(conn, dados).",
    ]
    for s in steps:
        doc.add_paragraph(s, style="List Number")
    doc.add_paragraph()
    add_para(doc, "A função inserir_ou_atualizar usa INSERT ... ON CONFLICT (orgao, tipo, ano, numero, origem) DO UPDATE, ou seja: se já existir um registro para a mesma chave natural, atualiza data_publicacao, titulo, ementa, url_oficial, caminho_html, caminho_html_bruto, caminho_docx e texto_busca, preservando os demais campos quando não fornecidos (COALESCE). Assim, reexecutar o coletor não duplica registros e pode atualizar conteúdo já existente.")
    doc.add_paragraph()

    add_heading(doc, "2.3 Outras fontes e verificação", level=2)
    add_para(doc, "Outras origens (govbr, dou, dodf, ocr, manual) podem ser alimentadas por outros scripts. O fluxo de verificação de normas (admin) atualiza os campos verificacao_status, verificacao_descricao, verificacao_em e verificacao_admin_id na tabela norma quando um administrador aprova ou rejeita a verificação de um ato.")
    doc.add_paragraph()

    add_heading(doc, "2.4 Variáveis de ambiente do banco", level=2)
    add_para(doc, "Os scripts e a API utilizam as variáveis DB_HOST, DB_PORT, DB_NAME, DB_USER e DB_PASSWORD (e opcionalmente DB_SSLMODE) para conectar ao PostgreSQL. O módulo db/db.py centraliza a conexão; scripts como coletar_arespcj.py e criar_admin.py devem ser executados com essas variáveis definidas (por exemplo na VPS).")
    doc.add_paragraph()

    doc.add_paragraph()
    add_para(doc, "— Documento gerado pelo script scripts/gerar_doc_banco.py —")

    doc.save(str(out_path))
    print(f"Documento salvo em: {out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
